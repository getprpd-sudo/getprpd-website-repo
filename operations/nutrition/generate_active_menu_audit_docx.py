"""Generate the durable Word version of the PRPD active-menu nutrition audit.

Design preset: compact_reference_guide.
Named override: PRPD forest-green brand accent replaces the preset blue accent.
Header pattern: memo_masthead.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import calculate_active_menu as audit


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT.parent / "PRPD_ACTIVE_MENU_NUTRITION_AUDIT.docx"

GREEN = "1E2E1E"
SAGE = "7FA882"
CREAM = "ECE7DF"
LIGHT = "F5F3EE"
RED = "8E2D2D"
GRAY = "666666"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(table) -> None:
    repeat_header(table.rows[0])
    for row in table.rows:
        prevent_row_split(row)


def set_paragraph_border_bottom(paragraph, color=GREEN, size=12, space=5) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(GREEN)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, 18, 10),
        "Heading 2": (13, 14, 7),
        "Heading 3": (12, 10, 5),
    }
    for style_name, (size, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(GREEN)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = "PRPD  |  ACTIVE MENU NUTRITION AUDIT"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.runs[0]
    header_run.font.name = "Calibri"
    header_run.font.size = Pt(8.5)
    header_run.font.bold = True
    header_run.font.color.rgb = rgb(SAGE)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("Internal working record  |  Page ")
    footer_run.font.name = "Calibri"
    footer_run.font.size = Pt(8.5)
    footer_run.font.color.rgb = rgb(GRAY)
    add_field(footer, "PAGE")


def add_title_block(doc: Document) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("RECIPE + NUTRITION CONTROL")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = rgb(SAGE)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("PRPD Active Menu Nutrition Audit")
    run.font.name = "Calibri"
    run.font.size = Pt(25)
    run.font.bold = True
    run.font.color.rgb = rgb(GREEN)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle_run = subtitle.add_run("Independent calorie and nutrient rebuild for the July 2026 menu")
    subtitle_run.font.name = "Calibri"
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = rgb(GRAY)
    set_paragraph_border_bottom(subtitle)

    metadata = [
        ("Audit date", "July 14, 2026"),
        ("Scope", "15 active dishes"),
        ("Status", "Complete calculated estimate"),
        ("Control rule", "Website values are comparison only"),
    ]
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        label_run.font.color.rgb = rgb(GREEN)
        value_run = paragraph.add_run(value)
        value_run.font.color.rgb = rgb(GRAY)


def add_callout(doc: Document, label: str, text: str, hold=False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(10)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F8EAEA" if hold else CREAM)
    p_pr.append(shd)
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.color.rgb = rgb(RED if hold else GREEN)
    body_run = paragraph.add_run(text)
    body_run.font.color.rgb = rgb(GREEN)


def fill_table_header(table, headers: list[str]) -> None:
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        set_cell_shading(cell, GREEN)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = rgb(WHITE)
                run.font.size = Pt(9)


def add_executive_results(doc: Document) -> None:
    doc.add_heading("Executive Results", level=1)
    doc.add_paragraph(
        "Every result below was recalculated from recorded ingredient quantities. Package labels and manufacturer data were used for known products; USDA records were used for whole foods. Assumptions are visible instead of being hidden inside a final-looking number."
    )
    add_callout(
        doc,
        "Bottom line",
        "All required Nutrition Facts nutrients now have calculated estimates. Use these values in PRPD's estimated-label workflow and regenerate them whenever a recipe, portion, or product changes.",
    )

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    fill_table_header(table, ["Dish", "Tier", "Audited", "Live", "Delta", "Decision"])
    for meal in audit.MEALS:
        for tier_name, build_record, current in audit.tier_records(meal):
            if build_record is None or current is None:
                continue
            value = audit.total(build_record)
            rounded_cal = audit.display_value(value.calories, "calories")
            rounded_protein = audit.display_value(value.protein, "macro")
            row = table.add_row()
            values = [
                meal.name,
                tier_name,
                f"{rounded_cal} kcal / {rounded_protein}g P",
                f"{current[0]} kcal / {current[1]}g P",
                f"{rounded_cal - current[0]:+d} kcal\n{rounded_protein - current[1]:+d}g P",
                "CHECK NEXT COOK" if build_record.confidence == "low" else "ESTIMATE READY",
            ]
            for index, value_text in enumerate(values):
                row.cells[index].text = value_text
                if index in (1, 4, 5):
                    for paragraph in row.cells[index].paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if build_record.confidence == "low":
                set_cell_shading(row.cells[5], "F8EAEA")
                for paragraph in row.cells[5].paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = rgb(RED)
                        run.bold = True
            else:
                set_cell_shading(row.cells[5], LIGHT)
    set_table_geometry(table, [2450, 700, 1450, 1450, 1250, 2060])
    set_repeat_table_header(table)


def add_scrutiny_checks(doc: Document) -> None:
    doc.add_heading("Calorie Scrutiny Checks", level=1)
    checks = [
        "Oreo Mousse: the confirmed 19.5g Premier Protein portion produces about 292 calories and 38g protein and is now locked in the active recipe.",
        "Bulk Seekh: the approved reduction to 210g raw seekh mixture across 1.5 breads rebuilds to about 770 calories and 56g protein while preserving the three-half presentation.",
        "French Toast: calories are close to the website after using the confirmed 80-calorie bread and 12g sugar. Protein is about 34g Lean and 45g Bulk, lower than published.",
        "Ground beef: the working model uses USDA low-fat ground-beef yields of 72% baked and 69% pan-browned until PRPD records its own cooked yield.",
        "Sodium: every savory serving includes a standardized 0.5g-1g salt estimate. Replace it with weighed batch salt divided by actual yield after the next cook.",
    ]
    for check in checks:
        doc.add_paragraph(check, style="List Bullet")


def format_amount(portion: audit.Portion) -> str:
    amount = f"{portion.amount:g} {portion.unit}"
    if portion.note:
        amount += f" ({portion.note})"
    return amount


def add_meal_detail(doc: Document, meal: audit.Meal) -> None:
    doc.add_heading(f"{meal.meal_id.upper()}  {meal.name}", level=1)
    doc.add_paragraph(f"Recipe status: {meal.status}")

    for tier_name, build_record, current in audit.tier_records(meal):
        if build_record is None:
            continue
        value = audit.total(build_record)
        doc.add_heading(tier_name, level=2)

        totals = doc.add_table(rows=2, cols=5)
        totals.style = "Table Grid"
        fill_table_header(totals, ["Calories", "Protein", "Carbs", "Fiber", "Fat"])
        summary_values = [
            f"{value.calories:.1f}",
            f"{value.protein:.1f}g",
            f"{value.carbs:.1f}g",
            f"{value.fiber:.1f}g",
            f"{value.fat:.1f}g",
        ]
        for cell, summary_value in zip(totals.rows[1].cells, summary_values):
            cell.text = summary_value
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_shading(cell, LIGHT)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        set_table_geometry(totals, [1872] * 5)
        set_repeat_table_header(totals)

        details = doc.add_paragraph()
        details.paragraph_format.space_before = Pt(4)
        details.add_run("Additional working nutrients: ").bold = True
        details.add_run(
            f"sat fat {value.saturated_fat:.1f}g; trans fat {value.trans_fat:.1f}g; sugars {value.sugars:.1f}g; added sugars {value.added_sugars:.1f}g; sodium {value.sodium:.0f}mg; cholesterol {value.cholesterol:.0f}mg; vitamin D {value.vitamin_d:.1f}mcg; calcium {value.calcium:.0f}mg; iron {value.iron:.1f}mg; potassium {value.potassium:.0f}mg."
        )
        if current:
            comparison = doc.add_paragraph()
            comparison.add_run("Live-site comparison: ").bold = True
            comparison.add_run(
                f"{current[0]} calories / {current[1]}g protein. Working delta: "
                f"{audit.display_value(value.calories, 'calories') - current[0]:+d} calories and "
                f"{audit.display_value(value.protein, 'macro') - current[1]:+d}g protein."
            )

        table = doc.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        fill_table_header(table, ["Ingredient", "Amount", "kcal", "P", "C", "F", "Source"])
        for portion in build_record.portions:
            ingredient = audit.INGREDIENTS[portion.ingredient]
            part = audit.portion_nutrition(portion)
            row = table.add_row()
            values = [
                ingredient.name,
                format_amount(portion),
                f"{part.calories:.1f}",
                f"{part.protein:.1f}",
                f"{part.carbs:.1f}",
                f"{part.fat:.1f}",
                ingredient.confidence.title(),
            ]
            for index, value_text in enumerate(values):
                row.cells[index].text = value_text
                if index >= 2:
                    row.cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        total_row = table.add_row()
        total_values = ["Calculated total", "", f"{value.calories:.1f}", f"{value.protein:.1f}", f"{value.carbs:.1f}", f"{value.fat:.1f}", build_record.confidence.title()]
        for index, value_text in enumerate(total_values):
            total_row.cells[index].text = value_text
            set_cell_shading(total_row.cells[index], CREAM)
            for run in total_row.cells[index].paragraphs[0].runs:
                run.bold = True
        set_table_geometry(table, [2050, 2350, 700, 650, 650, 650, 2310])
        set_repeat_table_header(table)

        doc.add_paragraph("Assumptions", style="Heading 3")
        for assumption in build_record.assumptions:
            doc.add_paragraph(assumption, style="List Bullet")

    add_callout(
        doc,
        "Recommendation",
        meal.recommendation,
        hold=any(b and b.confidence == "low" for b in (meal.lean, meal.bulk)),
    )


def add_critical_inputs(doc: Document) -> None:
    doc.add_heading("Next-Cook Accuracy Upgrades", level=1)
    items = [
        "Record rotating sauce-cup grams and replace generic package estimates only when a purchased product changes.",
        "Retained cooking oil or spray for egg bites, hash, French Toast, potatoes, broccoli, chicken, and corn.",
        "Butter Chicken sauce batch yield and the number of meals served by the listed formula.",
        "Peri Peri cooked edible meat weight and retained marinade from one Lean and one Bulk serving.",
        "Portion Chocolate Oreo Mousse with 19.5g Premier Protein powder and record finished cup weight when convenient.",
        "Current ladyfinger, mascarpone, Philadelphia filling, and shawarma-bread labels.",
        "Measured salt in each batch and the finished yield used to replace the standardized salt estimate.",
        "Finished serving/net weight for every dish and tier.",
    ]
    for index, item in enumerate(items, start=1):
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_label_field_status(doc: Document) -> None:
    doc.add_heading("Nutrition Facts Field Status", level=1)
    doc.add_paragraph(
        "The calculator covers every nutrient required on the standard U.S. Nutrition Facts panel. Exact package values are used where available and approved generic/database estimates are used for remaining inputs."
    )
    rows = [
        ("Calories", "Complete estimate", "Regenerate after recipe, portion, or product changes."),
        ("Protein, carbohydrate, fiber, fat", "Complete estimate", "Package values and USDA data are used."),
        ("Saturated fat, trans fat, sugars, added sugars, cholesterol", "Complete estimate", "Package values and generic equivalents are documented."),
        ("Sodium", "Complete estimate", "Includes standardized added-salt estimates; weigh salt by batch later."),
        ("Vitamin D, calcium, iron, potassium", "Complete estimate", "Package values and USDA data are used."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    fill_table_header(table, ["Field", "Status", "Remaining work"])
    for field_name, status, work in rows:
        row = table.add_row()
        row.cells[0].text = field_name
        row.cells[1].text = status
        row.cells[2].text = work
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if status == "Pending":
            set_cell_shading(row.cells[1], "F8EAEA")
            for run in row.cells[1].paragraphs[0].runs:
                run.font.color.rgb = rgb(RED)
                run.bold = True
        elif status == "Partial":
            set_cell_shading(row.cells[1], CREAM)
    set_table_geometry(table, [2500, 1200, 5660])
    set_repeat_table_header(table)


def add_sources(doc: Document) -> None:
    doc.add_heading("Method and Sources", level=1)
    doc.add_paragraph(
        "This is a database calculation, not laboratory analysis. Exact current package labels take priority over generic values. USDA cooked yields are temporary substitutes for missing kitchen yields. FDA-style declaration and rounding should occur only after serving weights and recipes are locked."
    )
    sources = [
        "USDA FoodData Central - fdc.nal.usda.gov",
        "USDA Cooking Yield Data for Meat and Poultry - low-fat ground-beef and meat-yield references",
        "FAGE Total 0% manufacturer nutrition page",
        "Mission Carb Balance manufacturer nutrition pages",
        "OREO Thins manufacturer nutrition page",
        "HERSHEY'S Natural Unsweetened Cocoa SmartLabel",
        "Premier Protein user-supplied package label (39g serving)",
        "Fairlife fat-free ultra-filtered milk manufacturer page",
        "H-E-B fat-free cottage cheese and mozzarella manufacturer pages",
        "Simple Truth Organic Plain Nonfat Greek Yogurt product nutrition page",
        "G Hughes Sugar Free Sweet Chili Sauce product nutrition page",
        "Great Value Light Mayonnaise product nutrition page",
        "FDA Guide for Developing and Using Databases for Nutrition Labeling",
        "FDA Daily Value and required nutrient guidance for Nutrition Facts labels",
    ]
    for source in sources:
        doc.add_paragraph(source, style="List Bullet")


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_title_block(doc)
    add_executive_results(doc)
    add_scrutiny_checks(doc)
    add_label_field_status(doc)
    for meal in audit.MEALS:
        add_meal_detail(doc, meal)
    add_critical_inputs(doc)
    add_sources(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
